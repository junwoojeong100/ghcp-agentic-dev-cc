#!/usr/bin/env python3
"""Build the editable, local Korean CXO presentation with python-pptx==1.0.2.

Usage (after the main workflow has produced the actual facts and screenshots):
    python3 scripts/build-deck.py
    python3 scripts/build-deck.py --font "Malgun Gothic" --require-screenshots

No network, application execution, or publication is performed. Verification
counts and implementation/review status come only from production/facts.json.
Missing screenshots are explicitly labeled placeholders unless strict mode is on.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    import pptx
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from PIL import Image
except ImportError as exc:
    raise SystemExit(
        'Required library missing. Install with: '
        'python3 -m pip install "python-pptx==1.0.2"'
    ) from exc

ROOT = Path(__file__).resolve().parents[1]
PRODUCTION = ROOT / "production"
OUTPUT = ROOT / "deliverables" / "ghcp-agentic-development.pptx"
WIDTH, HEIGHT = 13.333333, 7.5
COLORS = {
    "ground": "F3F6FA",
    "ink": "172B45",
    "accent": "006D77",
    "secondary": "52647A",
    "white": "FFFFFF",
    "line": "CDDBE7",
    "soft": "E0EFF1",
    "blocked": "9D3543",
    "blocked_bg": "FCECEF",
    "allowed": "17664C",
    "allowed_bg": "E9F4EE",
    "pending": "855811",
    "pending_bg": "FFF3DA",
}
SOURCES = (
    (
        "Copilot CLI 사용자 지정 에이전트 만들기",
        "https://docs.github.com/en/copilot/how-tos/copilot-cli/"
        "customize-copilot/create-custom-agents-for-cli",
    ),
    (
        "Copilot CLI 사용자 지정 에이전트 호출",
        "https://docs.github.com/en/copilot/how-tos/copilot-cli/"
        "use-copilot-cli/invoke-custom-agents",
    ),
    (
        "Copilot CLI 도구 실행 허용",
        "https://docs.github.com/en/copilot/how-tos/copilot-cli/"
        "use-copilot-cli/allowing-tools",
    ),
    (
        "Copilot Chat의 책임 있는 사용",
        "https://docs.github.com/en/copilot/responsible-use/copilot-chat",
    ),
)


def load_facts(path: Path) -> dict:
    if not path.is_file():
        raise ValueError(
            f"실행 근거가 없습니다: {path}\n"
            "주 작업에서 실제 구현·테스트·브라우저 검증을 마친 뒤 facts.json을 작성하세요. "
            "추정 수치로 슬라이드를 만들지 않습니다."
        )
    try:
        facts = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"facts.json을 읽을 수 없습니다: {exc}") from exc
    if not isinstance(facts, dict):
        raise ValueError("facts.json 최상위 값은 객체여야 합니다.")
    for key in ("copilotImplementation", "copilotReview"):
        if type(facts.get(key)) is not bool:
            raise ValueError(f"facts.json의 {key}는 명시적인 boolean이어야 합니다.")
    for parent, keys in (
        ("tests", ("passed", "failed")),
        ("browserChecks", ("passed",)),
        ("baselineTests", ("passed",)),
    ):
        value = facts.get(parent)
        if not isinstance(value, dict):
            raise ValueError(f"facts.json의 {parent} 객체가 필요합니다.")
        for key in keys:
            if type(value.get(key)) is not int or value[key] < 0:
                raise ValueError(f"{parent}.{key}는 0 이상의 정수여야 합니다.")
    for key in ("codeHash", "reviewStatus", "evidenceDir"):
        if not isinstance(facts.get(key), str) or not facts[key].strip():
            raise ValueError(f"facts.json의 {key} 문자열이 필요합니다.")
    limitations = facts.get("limitations")
    if not isinstance(limitations, list) or any(
        not isinstance(item, str) or not item.strip() for item in limitations
    ):
        raise ValueError("facts.json의 limitations는 문자열 배열이어야 합니다.")
    # A final deck must not claim successful behavior merely from prepared images.
    if not facts["copilotImplementation"]:
        raise ValueError("실제 Copilot 구현이 확인되지 않았습니다. 결과 시연 자료 생성을 보류합니다.")
    if facts["tests"]["failed"] or facts["tests"]["passed"] == 0:
        raise ValueError("구현 테스트의 성공 근거가 불충분합니다. 결과 시연 자료 생성을 보류합니다.")
    if facts["browserChecks"]["passed"] == 0 or facts["baselineTests"]["passed"] == 0:
        raise ValueError("시작 앱 또는 브라우저·API 검증 근거가 없습니다. 자료 생성을 보류합니다.")
    evidence = Path(facts["evidenceDir"])
    if not evidence.is_absolute():
        evidence = ROOT / evidence
    try:
        relative = evidence.resolve().relative_to(ROOT.resolve())
    except ValueError as exc:
        raise ValueError("evidenceDir은 이 프로젝트 내부 경로여야 합니다.") from exc
    if not evidence.is_dir():
        raise ValueError(f"실행 근거 디렉터리가 없습니다: {evidence}")
    facts["_evidenceLabel"] = relative.as_posix()
    return facts


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(COLORS.get(name, name))


class Deck:
    def __init__(self, facts: dict, font: str, require_screenshots: bool):
        self.facts = facts
        self.font = font
        self.require_screenshots = require_screenshots
        self.missing = []
        self.prs = Presentation()
        self.prs.slide_width = Inches(WIDTH)
        self.prs.slide_height = Inches(HEIGHT)
        self.prs.core_properties.title = "GitHub Copilot · Agentic Development in Action"
        self.prs.core_properties.subject = "기존 견적 앱에 최저 매출총이익률 규칙을 추가한 로컬 데모"
        self.prs.core_properties.author = "Demo production"
        self.prs.core_properties.keywords = "합성 데이터, 로컬 데모, 편집 가능한 한국어 슬라이드"
        self.prs.core_properties.comments = "공식 제품 소개자료나 공급자의 보증·후원을 나타내지 않습니다."

    def box(self, slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
        obj = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        obj.fill.solid()
        obj.fill.fore_color.rgb = rgb(fill)
        if line:
            obj.line.color.rgb = rgb(line)
            obj.line.width = Pt(0.8)
        else:
            obj.line.fill.background()
        return obj

    def text(self, slide, value, x, y, w, h, size=22, color="ink", bold=False,
             align=PP_ALIGN.LEFT, font=None, link=None):
        obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        obj.text_frame.clear()
        tf = obj.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 1.16
            p.space_before = Pt(0)
            p.space_after = Pt(3)
            run = p.add_run()
            run.text = line
            run.font.name = font or self.font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", "ko-KR")
            # Set East Asian typeface explicitly; Latin font alone is insufficient.
            for tag in ("a:ea", "a:cs"):
                element = rpr.find(tag, rpr.nsmap)
                if element is None:
                    element = OxmlElement(tag)
                    rpr.append(element)
                element.set("typeface", self.font)
        if link:
            # Keep the requested text color in PDF export. Run-level links are
            # forced to theme blue by Impress, unreadable on the video panel.
            obj.click_action.hyperlink.address = link
        return obj

    def notes(self, slide, speech: str, detail: str = ""):
        review = "실행 확인" if self.facts["copilotReview"] else "미확인"
        shared = (
            "제작 및 발표 범위\n"
            "합성 견적 데이터, localhost, 모의 전송만 사용한다. 운영 배포 및 외부 발송은 없다. "
            "실행 중 견적 판정은 결정적 서버 코드이며 견적마다 LLM을 호출하지 않는다. "
            "실제 시연 도구는 Copilot CLI다. VS Code UI, Copilot Studio, A2A 프로토콜은 "
            "이번 결과물에서 구현하거나 검증하지 않았다. 공식 소개자료·보증·후원을 뜻하지 않는다.\n"
            "H1은 구현 전에 범위를 확인하는 지점이다. Copilot의 node --test 실행은 권한 거절로 "
            "차단되었으며 Copilot이 테스트를 실행하거나 통과시킨 것이 아니다. 사용자 별도 승인에 "
            "따른 로컬 검증을 거친 뒤 Copilot 읽기 전용 검토로 이어진다. 실제 승인 채팅이나 "
            "버튼을 재현하지 않는다. H2 결과 채택·출시 판단은 아직 사람에게 남아 있다.\n"
            "영상은 실제 앱과 실행 근거를 편집한 사전 제작물이며 한국어 합성 음성·자막을 사용한다. "
            "영상 길이는 개발 소요시간 또는 생산성 향상 근거가 아니다.\n"
            f"글꼴: 기본 {self.font}. macOS Apple SD Gothic Neo, Windows Malgun Gothic "
            "(맑은 고딕) 대체. 글꼴 미포함. 대체 시 줄바꿈을 확인하거나 --font 옵션으로 다시 만든다. "
            "본문·도형은 편집 가능하며 실제 앱 캡처만 이미지다.\n"
            f"실행 근거: {self.facts['_evidenceLabel']}\n"
            f"검증 코드 식별자: {self.facts['codeHash']}\n"
            f"Copilot 검토 실행: {review}; 기록 상태: {self.facts['reviewStatus']}\n"
            "기록된 제약:\n" + "\n".join(f"- {item}" for item in self.facts["limitations"])
        )
        slide.notes_slide.notes_text_frame.text = (
            "발표자 메모\n" + speech + ("\n\n" + detail if detail else "") + "\n\n" + shared
        )

    def slide(self, section, title=None, subtitle=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb("ground")
        self.text(slide, "GitHub Copilot", .55, .23, 4, .3, 13, bold=True, font="Aptos")
        self.text(slide, section, 7.1, .23, 5.68, .3, 12, "secondary", align=PP_ALIGN.RIGHT)
        self.box(slide, .55, .69, 12.23, .012, "line")
        if title:
            self.text(slide, title, .55, .94, 12.23, 1.03, 32, bold=True)
        if subtitle:
            self.text(slide, subtitle, .57, 1.83, 12.1, .48, 20, "secondary")
        self.box(slide, .55, 6.99, 12.23, .012, "line")
        self.text(slide, "사전 제작 데모 · 합성 데이터 · 모의 전송 · 운영 배포 없음", .55, 7.12, 9.8, .22, 10, "secondary")
        self.text(slide, f"{len(self.prs.slides):02d} / 12", 11.88, 7.1, .9, .25, 11, "secondary", align=PP_ALIGN.RIGHT)
        return slide

    def status(self, slide, label, x, y, w, kind):
        self.box(slide, x, y, w, .44, f"{kind}_bg")
        self.text(slide, label, x + .12, y + .055, w - .24, .34, 17, kind, bold=True)

    def screenshot(self, slide, name, x, y, w, h):
        path = PRODUCTION / "screenshots" / f"{name}.png"
        self.box(slide, x, y, w, h, "white", "line")
        if not path.is_file():
            if self.require_screenshots:
                raise ValueError(f"실제 캡처가 없습니다: {path}")
            self.missing.append(path)
            self.text(slide, "실제 캡처 준비 중", x + .25, y + h / 2 - .3, w - .5, .55,
                      24, "secondary", align=PP_ALIGN.CENTER)
            self.text(slide, path.name, x + .2, y + h / 2 + .32, w - .4, .3,
                      12, "secondary", align=PP_ALIGN.CENTER, font="Aptos")
            return
        with Image.open(path) as image:
            iw, ih = image.size
            image.verify()
        if iw <= 0 or ih <= 0:
            raise ValueError(f"캡처의 크기가 올바르지 않습니다: {path}")
        pad = .035
        scale = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
        pw, ph = iw * scale, ih * scale
        picture = slide.shapes.add_picture(
            str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
            width=Inches(pw), height=Inches(ph),
        )
        for props in picture._element.xpath(".//p:cNvPr"):
            props.set("descr", f"실제 로컬 앱 캡처 {name}. 출처 production/screenshots/{name}.png")

    def build(self):
        self.title()
        self.connected()
        self.business()
        self.scope()
        self.workflow()
        self.video()
        self.result()
        self.evidence()
        self.boundaries()
        self.pilot()
        self.closing()
        self.sources()
        assert len(self.prs.slides) == 12
        for number, slide in enumerate(self.prs.slides, 1):
            for shape in slide.shapes:
                if (shape.left < 0 or shape.top < 0 or
                        shape.left + shape.width > self.prs.slide_width + 10 or
                        shape.top + shape.height > self.prs.slide_height + 10):
                    raise ValueError(f"슬라이드 {number}: 슬라이드 밖으로 나간 개체 {shape.name}")
        return self.prs

    def title(self):
        s = self.slide("Agentic Development in Action")
        self.text(s, "GitHub Copilot", .57, 1.44, 11.8, .62, 34, "accent", True, font="Aptos")
        self.text(s, "사업 요청이 제품의 동작이 되는 순간", .55, 2.4, 12.15, 1.08, 34, bold=True)
        self.text(s, "Agentic Development in Action", .58, 3.65, 11.9, .64, 25, "secondary", font="Aptos")
        self.text(s, "이 견적, 보내도 됩니까?", .58, 5.05, 11.9, .61, 26, bold=True)
        self.text(s, "최저 매출총이익률 15% 규칙을 기존 앱에 연결한 작은 변경", .58, 5.78, 11.9, .63, 21, "secondary")
        self.notes(s,
            "오늘의 질문은 이 견적을 보내도 되는가입니다. 거대한 새 시스템을 제안하는 대신, "
            "기존 견적 앱의 전송 규칙 하나가 사업 요청에서 코드와 실행 근거로 바뀌는 모습을 보겠습니다. "
            "GitHub Copilot CLI를 사용한 로컬 사례이며 공식 제품 소개자료가 아닙니다.")

    def connected(self):
        s = self.slide("왜 연결된 개발인가", "답변에서, 확인할 수 있는 변경으로")
        self.text(s, "대화에 머무는 요청", .6, 2.05, 5.5, .5, 24, bold=True)
        self.box(s, .6, 2.81, 5.49, 1.47, "white", "line")
        self.text(s, "15% 미만 견적은\n전송을 막아 주세요.", .87, 3.05, 4.94, 1.09, 25)
        self.text(s, "요청 문구 예시 · 실제 사용자 대화 재현 아님", .64, 4.49, 5.5, .45, 13, "secondary")
        self.box(s, 6.43, 2.05, .015, 4.4, "line")
        self.text(s, "코드에 연결된 요청", 6.91, 2.05, 5.8, .5, 24, "accent", True)
        rows = (("현재 파일", "기존 계산·화면·전송 처리 확인"),
                ("실제 변경", "차단 규칙과 안내를 코드에 반영"),
                ("실행 근거", "별도 로컬 테스트·브라우저·API 기록"))
        for i, (label, value) in enumerate(rows):
            y = 2.94 + i * 1.03
            self.text(s, label, 6.93, y, 5.6, .4, 20, "accent", True)
            self.text(s, value, 6.93, y + .42, 5.75, .48, 20)
        self.text(s, "이번 데모는 로컬 파일과 실행 결과를 연결합니다.", .6, 6.15, 11.9, .5, 21, "secondary")
        self.notes(s,
            "왼쪽은 업무 요청 문구를 읽기 쉽게 따로 보여 준 예시입니다. 실제 승인 대화가 아닙니다. "
            "오른쪽은 Copilot이 현재 파일을 읽고 코드를 바꾼 뒤, 사용자 승인에 따른 별도 로컬 "
            "검증이 실행 근거를 남기는 흐름입니다. Copilot이 테스트를 실행했다는 뜻이 아닙니다. "
            "다른 도구가 할 수 없다는 주장이나 독점 기능 비교는 하지 않습니다.")

    def business(self):
        s = self.slide("업무 요청", "10% 할인 견적은 왜 멈춰야 할까요?")
        self.status(s, "변경 전 · 차단 규칙 없음", .58, 1.96, 3.6, "pending")
        self.screenshot(s, "before", .58, 2.6, 7.73, 4.27)
        self.text(s, "준비된 합성 견적", 8.65, 2.04, 4.08, .49, 24, bold=True)
        values = (
            "판매가 1,000만 원",
            "등록 원가 800만 원",
            "10% 할인 → 견적액 900만 원",
            "이익 100만 원",
        )
        for i, value in enumerate(values):
            self.text(s, value, 8.65, 2.77 + i * .59, 4.12, .51, 20)
        self.box(s, 8.65, 5.34, 4.08, .012, "line")
        self.text(s, "100만 ÷ 900만 ≈ 11.1%", 8.65, 5.59, 4.12, .6, 22, "blocked", True)
        self.text(s, "최저 기준 15%에 미달", 8.65, 6.24, 4.12, .48, 20, "blocked")
        self.notes(s,
            "여기서 이익률의 분모는 할인 전 판매가가 아니라 할인 후 견적액입니다. "
            "900만 원을 받고 원가 800만 원을 빼면 이익은 100만 원, 매출총이익률은 약 11.1%입니다. "
            "시작 앱은 계산과 모의 전송을 위해 미리 준비한 합성 앱이며 Copilot이 처음부터 "
            "작성한 것처럼 소개하지 않습니다. 이후 추가한 기능은 최소 기준 미달 전송 차단입니다.")

    def scope(self):
        s = self.slide("한 가지 변경", "규칙은 작게, 판정은 정확하게")
        self.text(s, "15% 미만이면 차단", .59, 2.03, 6.11, .63, 28, "blocked", True)
        self.text(s, "정확히 15%면 허용", .59, 2.91, 6.11, .63, 28, "allowed", True)
        self.text(s, "화면에서 이유를 안내하고\n서버에서도 같은 기준을 검사합니다.", .6, 3.91, 6.0, 1.07, 22)
        self.text(s, "표시용 반올림값으로 판정하지 않습니다.", .6, 5.42, 6.05, .79, 21, "secondary")
        self.status(s, "경계값 · 원가 1원 차이도 검사", 6.91, 2.03, 5.81, "pending")
        self.screenshot(s, "boundary", 6.91, 2.74, 5.81, 3.37)
        self.text(s, "견적액 1,000만 원 · 원가 8,500,001원 → 미달", 6.91, 6.28, 5.81, .61, 20)
        self.notes(s,
            "기준은 미만입니다. 정확히 15%인 견적은 허용해야 정상 업무를 불필요하게 막지 않습니다. "
            "또 1,000만 원 견적에서 원가가 850만 1원이면 실제 이익률은 15%에 조금 못 미칩니다. "
            "화면이 15.0%로 반올림해 보여 주더라도 차단하는 사례입니다. "
            "여기에 승인 라인, 예외 결재, 외부 발송 시스템을 추가한 것은 아닙니다.")

    def workflow(self):
        s = self.slide("역할과 전달 근거", "권한의 경계를 지키며, 실행 근거를 전달")
        self.text(s, "H1  사람 · 구현 전 범위 확인", .61, 1.99, 12.04, .59, 23, "accent", True)
        self.text(s, "Copilot 테스트 권한 거절 → 사용자 별도 승인 후 로컬 검증", .61, 2.61, 12.04, .51, 21, "secondary")
        roles = (("Copilot 구현", "수정 코드"), ("별도 로컬 검증", "테스트·UI·API"),
                 ("Copilot 검토", "읽기 전용 의견"), ("사람의 채택 판단", "아직 미결정"))
        for i, (role, artifact) in enumerate(roles):
            x = .63 + i * 3.14
            self.box(s, x, 3.66, 2.57, 1.39, "white", "line")
            self.text(s, role, x + .15, 3.87, 2.27, .5, 20, "accent", True)
            self.text(s, artifact, x + .15, 4.5, 2.27, .42, 19)
            if i < 3:
                self.box(s, x + 2.71, 4.26, .26, .27, "secondary", shape=MSO_SHAPE.CHEVRON)
                self.text(s, ("코드 전달", "결과 전달", "검토 전달")[i], x + 2.32,
                          3.22, 1.16, .34, 11, "secondary", align=PP_ALIGN.CENTER)
        if not self.facts["copilotReview"]:
            self.text(s, "Copilot 검토 실행은 미확인입니다.", 6.85, 5.18, 2.73, .75, 17, "pending")
        self.text(s, "H2  사람 · 결과 채택·출시 판단은 아직 남아 있습니다.", .61, 6.17, 12.05, .57, 23, bold=True)
        self.notes(s,
            "흐름은 Copilot 구현 → 별도 로컬 검증 → Copilot 읽기 전용 검토 → 사람의 채택 판단입니다. "
            "실제 Copilot 구현 중 node --test 실행은 권한 거절로 차단되었습니다. 자동으로 권한을 "
            "우회한 것이 아니라, 이후 사용자가 별도 로컬 테스트와 브라우저 검증을 승인한 것입니다. "
            "그 실행 근거를 읽기 전용 검토에 전달합니다. 구현 범위 확인과 로컬 검증 승인을 "
            "결과 채택이나 운영 배포 승인으로 바꾸어 말하지 않습니다. 검토도 독립 감사나 "
            "사람의 최종 승인이 아닙니다.",
            "계획의 업무 기준 출처: docs/change-request.md. 실제 Copilot 구현·검토 여부는 "
            "슬라이드 8의 facts.json 상태를 따른다. 사용자 지정 에이전트 설정이나 병렬 실행, "
            "A2A 통신을 이 그림만으로 구현했다고 주장하지 않는다.")

    def video(self):
        s = self.slide("5분 영상", "먼저 MP4 하나로 보여 주세요")
        self.box(s, .59, 2.12, 7.27, 3.84, "ink")
        self.text(s, "약 5분 · 한국어", .93, 2.67, 6.61, .64, 30, "white", True)
        self.text(s, "실제 앱 화면과 실행 기록을\n편집한 오프라인 데모", .93, 3.67, 6.55, 1.17, 25, "white")
        self.text(s, "ghcp-cxo-demo-ko.mp4", .94, 5.24, 6.56, .43, 21, "white", font="Aptos",
                  link="ghcp-cxo-demo-ko.mp4")
        self.text(s, "발표 순서", 8.38, 2.22, 4.3, .54, 25, bold=True)
        for i, value in enumerate(("1  MP4를 열고 전체 화면 재생", "2  필요할 때 멈춰 설명", "3  질문은 PPT와 근거로 확인")):
            self.text(s, value, 8.39, 3.12 + i * .89, 4.28, .68, 21)
        self.text(s, "합성 데이터·합성 음성 · 한국어 자막 · 대기 구간 편집 · 개발시간 보장 아님",
                  .61, 6.39, 12.1, .37, 13, "secondary")
        self.notes(s,
            "실전 발표에서는 PPT보다 MP4를 먼저 여는 편이 쉽습니다. 네트워크 없이 재생되는 "
            "사전 제작 데모이고, 이 슬라이드는 영상 재생을 안내하는 페이지입니다. "
            "MP4는 PPT 안에 삽입하지 않고 같은 deliverables 폴더의 별도 파일로 제공합니다. "
            "파일 링크가 재생되지 않으면 탐색기나 Finder에서 직접 열어 주세요. "
            "한국어 자막은 영상에 표시되며 SRT 파일도 별도로 제공합니다. 영상의 5분을 "
            "개발이 5분 만에 끝났다는 뜻으로 설명하지 않습니다.")

    def result(self):
        s = self.slide("같은 앱, 달라진 결과", "기준 미달은 차단하고, 정상 견적은 허용")
        self.status(s, "차단 · 할인 10% / 이익률 약 11.1%", .59, 1.99, 5.96, "blocked")
        self.status(s, "허용 · 할인 0% / 이익률 20%", 6.79, 1.99, 5.96, "allowed")
        self.screenshot(s, "blocked", .59, 2.67, 5.96, 3.46)
        self.screenshot(s, "allowed", 6.79, 2.67, 5.96, 3.46)
        self.text(s, "차단 이유 표시 · 모의 전송 불가", .6, 6.34, 5.95, .55, 21, "blocked", True)
        self.text(s, "정상 모의 전송 · 외부 발송 없음", 6.8, 6.34, 5.94, .55, 21, "allowed", True)
        self.notes(s,
            "왼쪽은 기존에 통과하던 10% 할인 견적입니다. 이제 기준 미달 이유를 보이고 "
            "전송을 막습니다. 오른쪽은 같은 견적의 할인을 없앤 경우입니다. 매출총이익률이 "
            "20%이므로 정상적으로 모의 전송됩니다. 색만으로 결론을 구분하지 않고 차단과 "
            "허용 문구를 함께 읽어 주세요. 이 캡처의 업무 동작을 영상에서도 확인할 수 있습니다.")

    def evidence(self):
        f = self.facts
        s = self.slide("실제 실행 근거", "완료라는 말보다, 남아 있는 근거")
        rows = (
            ("시작 앱 로컬 테스트", f"통과 {f['baselineTests']['passed']}건"),
            ("변경 후 로컬 테스트", f"통과 {f['tests']['passed']}건 · 실패 {f['tests']['failed']}건"),
            ("브라우저·API 확인", f"통과 {f['browserChecks']['passed']}건"),
            ("Copilot CLI", "구현 확인 · 읽기 전용 검토 " + ("실행 확인" if f["copilotReview"] else "미확인")),
        )
        self.text(s, "확인 항목", .65, 1.97, 4.0, .42, 20, "secondary", True)
        self.text(s, "production/facts.json 기록", 5.07, 1.97, 7.62, .42, 20, "secondary", True)
        for i, (label, value) in enumerate(rows):
            y = 2.6 + i * .7
            self.box(s, .61, y - .08, 12.08, .012, "line")
            self.text(s, label, .66, y + .05, 4.12, .51, 21)
            self.text(s, value, 5.07, y + .05, 7.63, .52, 21, "accent", True)
        self.box(s, .61, 5.48, 12.08, .012, "line")
        self.text(s, "테스트 실행: 별도 로컬 실행 (사용자 승인)\n"
                  "Copilot의 node --test는 권한 거절로 실행되지 않음",
                  .66, 5.65, 12.01, .7, 18, "ink", True)
        # Keep the full review status and limitations in notes; counts still
        # come only from facts.json, never from Copilot's attempted test command.
        review_status = " ".join(f["reviewStatus"].split())
        if len(review_status) > 32:
            review_status = review_status[:31] + "…"
        self.text(s, f"검토: {review_status} · 추가 제약 {len(f['limitations'])}건 · 전문은 발표자 노트",
                  .66, 6.43, 12.01, .3, 14, "secondary")
        code = f["codeHash"]
        short_code = code if len(code) <= 20 else code[:20] + "…"
        self.text(s, f"코드 식별자  {short_code}  ·  근거  {f['_evidenceLabel']}", .66, 6.76, 12.01, .21, 10, "secondary", font="Aptos")
        self.notes(s,
            "여기 수치는 주 작업이 별도 로컬 검증 후 기록한 facts.json에서 읽어 온 값입니다. "
            "Copilot은 구현 중 테스트 실행 권한을 받지 못했으므로 이 통과 건수를 Copilot이 "
            "실행한 테스트 결과라고 말하지 않습니다. 사용자 별도 승인으로 수행한 로컬 검증입니다. "
            "시작 앱과 변경 후 테스트는 별도 실행이므로 더해서 하나의 고유 테스트 수처럼 "
            "소개하지 않습니다. 브라우저·API 확인은 기록된 UI·직접 API 확인 건수이며 사용자 수나 "
            "성능 지표가 아닙니다. 읽기 전용 검토도 사람의 채택 승인은 아닙니다. "
            "발표 직전 코드가 바뀌었다면 이 수치를 그대로 재사용하지 말고 다시 검증해야 합니다.",
            f"근거 디렉터리: {f['_evidenceLabel']}\n"
            "implement.log: 실제 구현 및 Copilot 테스트 실행 권한 거절\n"
            "local-tests.log, local-test-record.json: 별도 로컬 테스트 출력·실행 기록\n"
            "browser-checks.json, api-responses.json: 브라우저·API 확인 기록\n"
            "review.log: Copilot 읽기 전용 검토 기록; 실행 여부는 facts.json을 따른다\n"
            "changes.diff: 실제 코드 변경\n"
            "제약은 이 슬라이드의 발표자 메모 아래에 facts.json 원문 항목으로 함께 수록한다. "
            "누락된 메타데이터나 실패 테스트가 있으면 생성기가 결과 시연 자료 생성을 중단한다.")

    def boundaries(self):
        s = self.slide("통제와 구분", "AI가 개발을 돕고, 서버 코드가 견적을 판정")
        self.text(s, "견적마다 LLM을 호출하지 않습니다. 정해진 규칙을 서버가 실행합니다.",
                  .6, 1.98, 12.13, .94, 23, bold=True)
        rows = (
            ("개발 워크플로", "Copilot 구현·읽기 전용 검토 / 별도 로컬 검증"),
            ("Copilot Studio", "이번 시연에 미포함 · 구현·검증하지 않음"),
            ("A2A 프로토콜", "이번 시연에 미포함 · 구현·검증하지 않음"),
        )
        for i, (label, value) in enumerate(rows):
            y = 3.23 + i * .89
            self.box(s, .61, y - .12, 12.11, .012, "line")
            self.text(s, label, .66, y, 3.06, .55, 22, "accent", True)
            self.text(s, value, 3.92, y, 8.75, .62, 21)
        self.text(s, "localhost만 사용 · 실제 발송 없음 · 운영 배포 없음", .62, 6.22, 12.1, .57, 22, "secondary")
        self.notes(s,
            "AI를 개발 과정에 사용하는 것과 AI가 매 견적의 전송을 결정하는 것은 다릅니다. "
            "이 앱은 후자가 아닙니다. 화면에서도 서버에서도 결정적 규칙의 결과를 사용합니다. "
            "아래 세 항목은 제품 기능 우열표가 아니라 이번 시연의 포함 범위를 나눈 것입니다. "
            "Copilot Studio 앱이나 A2A 기반 에이전트 간 통신을 구현했다고 설명하지 않습니다. "
            "VS Code 화면은 검증하지 않았으므로 CLI 실행 근거로만 이야기합니다.")

    def pilot(self):
        s = self.slide("첫 적용 후보", "작고 분명한 요청 하나부터 고르세요")
        rows = (
            ("범위", "변경할 기능 하나와 제외할 일을 명시한다"),
            ("기준", "정상·예외·경계값을 말로 설명할 수 있다"),
            ("검증", "완료 기준을 실행 가능한 테스트로 확인한다"),
            ("복구", "이전 동작으로 돌아갈 경로를 보관한다"),
            ("책임", "최종 채택을 판단할 검토 책임자를 정한다"),
        )
        for i, (label, value) in enumerate(rows):
            y = 2.05 + i * .78
            self.text(s, label, .62, y, 1.39, .57, 23, "accent", True)
            self.text(s, value, 2.2, y, 10.52, .57, 23)
            self.box(s, .63, y + .64, 12.05, .012, "line")
        self.text(s, "H2 결과 채택·출시 판단은 사람에게 남겨 둡니다.", .63, 6.28, 12.03, .54, 22, "secondary")
        self.notes(s,
            "다음 행동은 전사 도입을 한 번에 약속하는 것이 아니라 후보 요청 하나를 고르는 일입니다. "
            "예외가 무한한 업무보다 범위를 좁힐 수 있고 결과를 테스트할 수 있는 변경이 좋습니다. "
            "원본과 되돌아갈 경로를 보관하고, 테스트가 통과한 뒤에도 누가 채택을 판단할지 정해 주세요. "
            "이 사례만으로 비용 절감이나 작업 시간 단축 수치를 산출하지 않습니다.")

    def closing(self):
        s = self.slide("Agentic Development in Action")
        self.text(s, "GitHub Copilot", .61, 1.63, 12.02, .55, 28, "accent", True, font="Aptos")
        self.text(s, "사업의 결정을 코드로.\n변경의 근거를 테스트로.\n출시 결정은 사람에게.",
                  .6, 2.63, 12.06, 2.58, 34, bold=True)
        self.text(s, "다음에는 우리 조직의 작은 요청 하나로 확인해 보세요.", .63, 5.97, 12.05, .66, 22, "secondary")
        self.notes(s,
            "요청을 코드로 연결하고, 변경의 근거를 실행 결과로 남기며, 마지막 의사결정은 "
            "사람에게 남기는 것이 오늘의 핵심입니다. 질문을 받으면 실제 화면, 실행 근거, "
            "남아 있는 제약을 순서대로 확인해 주세요. 아직 남은 채택 판단을 이미 완료된 "
            "승인처럼 마무리하지 않습니다.")

    def sources(self):
        s = self.slide("공식 참고 문서", "개념의 출처와 실행 근거를 구분합니다")
        for i, (label, url) in enumerate(SOURCES):
            y = 1.96 + i * 1.02
            self.text(s, label, .63, y, 12.03, .41, 20, "accent", True, link=url)
            # Long source URLs are the one small-font appendix exception.
            split = url.replace("/customize-copilot/", "/customize-copilot/\n").replace(
                "/use-copilot-cli/", "/use-copilot-cli/\n")
            self.text(s, split, .65, y + .44, 12.0, .51, 12, "secondary", font="Aptos", link=url)
        self.box(s, .63, 6.2, 12.02, .012, "line")
        self.text(s, "본 데모의 실행 근거", .65, 6.37, 3.4, .35, 17, bold=True)
        self.text(s, "production/facts.json  ·  " + self.facts["_evidenceLabel"],
                  4.0, 6.39, 8.68, .42, 13, "secondary", font="Aptos")
        self.notes(s,
            "공식 문서는 Copilot CLI와 책임 있는 사용의 개념을 확인할 때 참고합니다. "
            "사용자 지정 에이전트 관련 링크를 실었다고 해서 해당 구성을 이번 시연에서 "
            "모두 구현한 것은 아닙니다. 실행 여부와 결과는 로컬 facts.json 및 연결된 로그가 "
            "근거입니다. 문서 링크의 제공 사실과 링크 내용의 최신 확인은 구분합니다.",
            "제공된 공식 참고 URL:\n" + "\n".join(url for _, url in SOURCES))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", choices=("Apple SD Gothic Neo", "Malgun Gothic"),
                        default="Apple SD Gothic Neo" if sys.platform == "darwin" else "Malgun Gothic")
    parser.add_argument("--require-screenshots", action="store_true",
                        help="Fail rather than include explicitly labeled missing-capture placeholders.")
    args = parser.parse_args()
    try:
        if pptx.__version__ != "1.0.2":
            raise ValueError(f"python-pptx==1.0.2가 필요합니다. 현재 버전: {pptx.__version__}")
        facts = load_facts(PRODUCTION / "facts.json")
        deck = Deck(facts, args.font, args.require_screenshots)
        presentation = deck.build()
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(OUTPUT)
        print(OUTPUT)
        print("12 editable slides; Korean presenter notes; no publication.")
        for path in sorted(set(deck.missing)):
            print(f"WARNING: labeled screenshot placeholder: {path}", file=sys.stderr)
        if not facts["copilotReview"]:
            print("WARNING: Copilot review is explicitly marked unverified.", file=sys.stderr)
        return 0
    except (OSError, ValueError) as exc:
        print(f"DECK NOT BUILT: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
