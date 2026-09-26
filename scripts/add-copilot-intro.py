#!/usr/bin/env python3
"""Prepend four sourced, editable Korean slides to a NEW copy of a 12-slide deck.

Usage:
  python3 -B scripts/add-copilot-intro.py \
    --input deliverables/ghcp-live/ghcp-agentic-development.pptx \
    --output deliverables/ghcp-live/with-intro/ghcp-agentic-development.pptx

No network, recording, app execution, evidence updates or publication occurs.
The input must have 12 slides; the output must not already exist. The original
12 slides retain their historical findings and limitations. Only page numbers,
relative external links and media-location guidance change in the copied slides.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit, urlunsplit
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "production/ghcp-live/copilot-intro.json"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PAGE = re.compile(r"^\d{2} / (12|16)$")


def helpers():
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location("_intro_helpers", ROOT / "scripts/build-deck.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_content():
    content = json.loads(CONTENT.read_text(encoding="utf-8"))
    if content["schemaVersion"] != 1 or len(content["slides"]) != 4:
        raise ValueError("Exactly four schemaVersion 1 introduction slides are required.")
    for slide in content["slides"]:
        if not slide["notes"].strip() or not slide["sourceIds"]:
            raise ValueError("Every introduction slide needs notes and official sources.")
        for key in slide["sourceIds"]:
            url = content["sources"][key]["url"]
            if urlsplit(url).hostname not in {"docs.github.com", "code.visualstudio.com", "learn.microsoft.com"}:
                raise ValueError(f"Unexpected source host: {url}")
    return content


def make_intro(base, presentation, content, font):
    class Intro(base.Deck):
        def __init__(self):
            self.prs = presentation
            self.font = font

        def frame(self, record):
            s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = base.rgb("ground")
            self.text(s, "개발 실행 사례 · 제품 이해", .55, .23, 5.3, .3, 13, bold=True)
            self.text(s, record["section"], 6.6, .23, 6.18, .3, 12, "secondary", align=base.PP_ALIGN.RIGHT)
            self.box(s, .55, .69, 12.23, .012, "line")
            self.text(s, record["title"], .55, .96, 12.23, .78, 30, bold=True)
            self.text(s, record["subtitle"], .57, 1.81, 12.1, .55, 18, "secondary")
            self.box(s, .55, 6.99, 12.23, .012, "line")
            self.text(s, f"비공식 제품 개념 소개 · 공식 문서 확인 {content['checkedDate']} · 출처·제약은 발표자 노트",
                      .55, 7.12, 10.95, .22, 10, "secondary")
            self.text(s, "00 / 16", 11.88, 7.1, .9, .25, 11, "secondary", align=base.PP_ALIGN.RIGHT)
            source_notes = "\n\n".join(
                f"[{key}] {content['sources'][key]['title']}\n{content['sources'][key]['url']}\n"
                f"확인 내용: {content['sources'][key]['verified']}"
                for key in record["sourceIds"]
            )
            s.notes_slide.notes_text_frame.text = (
                "발표자 메모\n" + record["notes"] + "\n\n공식 문서 확인일: " + content["checkedDate"] +
                "\n자료 범위: " + content["scope"] + "\n\n공식 출처\n" + source_notes +
                "\n\n제품 기능은 플랜·버전·조직 정책에 따라 달라집니다. 제품 문서는 이 데모의 실행 증거가 아닙니다. "
                "뒤에 이어지는 12장은 입력한 데모 자료의 실행 근거와 제약을 그대로 유지합니다. "
                "소개 슬라이드 추가만으로 데모의 검증 상태나 촬영 완료 여부가 바뀌지 않습니다. "
                "GitHub 또는 Microsoft의 공식 발표자료·보증·후원을 뜻하지 않습니다."
            )
            return s

        def build(self):
            a, b, c, d = content["slides"]
            s = self.frame(a)
            self.text(s, a["headline"], .62, 2.55, 11.98, 1.2, 29, "accent", True)
            for i, item in enumerate(a["items"]):
                x = .62 + i * 4.15
                self.box(s, x, 4.13, 3.82, 1.66, "white", "line")
                self.text(s, item["heading"], x + .2, 4.31, 3.42, .42, 21, "accent", True)
                self.text(s, item["body"], x + .2, 4.87, 3.42, .74, 20)
            self.text(s, a["takeaway"], .64, 6.28, 11.98, .48, 18, bold=True)

            s = self.frame(b)
            for i, item in enumerate(b["items"]):
                x = .62 + (i % 2) * 6.22
                y = 2.50 + (i // 2) * 1.92
                self.box(s, x, y, 5.88, 1.78, "soft" if i == 3 else "white", "line")
                self.text(s, item["heading"], x + .2, y + .13, 5.48, .35, 21, "accent", True)
                self.text(s, item["detail"], x + .2, y + .57, 5.48, .32, 15, "secondary")
                self.text(s, item["body"], x + .2, y + .94, 5.48, .75, 17)
            self.text(s, b["takeaway"], .64, 6.45, 11.98, .36, 17, "secondary")

            s = self.frame(c)
            self.text(s, c["localLabel"], .63, 2.5, 7.67, .45, 22, "accent", True)
            for i, item in enumerate(c["items"]):
                y = 3.16 + i * .58
                self.box(s, .63, y - .07, 7.72, .012, "line")
                self.text(s, item["heading"], .66, y, 1.32, .43, 20, "accent", True)
                self.text(s, item["body"], 2.05, y + .02, 6.23, .42, 18)
            self.text(s, c["availability"], .66, 5.86, 7.67, .79, 15, "secondary")
            self.box(s, 8.71, 2.52, 4.0, 4.12, "soft")
            self.text(s, c["cliLabel"], 8.96, 2.75, 3.5, .45, 23, "accent", True)
            for i, item in enumerate(c["cliItems"]):
                y = 3.43 + i * .77
                self.text(s, item["heading"], 8.98, y, 3.44, .35, 20, bold=True)
                self.text(s, item["body"], 8.98, y + .4, 3.44, .32, 17, "secondary")
            self.text(s, c["takeaway"], 8.98, 6.12, 3.44, .33, 17, "accent", True)

            s = self.frame(d)
            for i, item in enumerate(d["items"]):
                y = 2.6 + i * 1.13
                self.text(s, item["heading"], .64, y, 6.77, .39, 21, "accent", True)
                self.text(s, item["body"], .66, y + .45, 6.9, .68, 18)
            self.box(s, 7.94, 2.56, 4.77, 3.71, "white", "line")
            self.text(s, d["contrastTitle"], 8.17, 2.75, 4.31, .77, 20, "accent", True)
            for i, item in enumerate(d["contrastItems"]):
                y = 3.73 + i * .79
                self.text(s, item["heading"], 8.18, y, 4.29, .32, 17, bold=True)
                self.text(s, item["body"], 8.18, y + .36, 4.29, .57, 17)
            self.text(s, d["availability"], 8.18, 5.66, 4.29, .53, 14, "secondary")
            self.text(s, d["takeaway"], .65, 6.46, 12.0, .36, 17, "secondary")
    Intro().build()


def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from all_shapes(shape.shapes)


def rebase_links(blob, input_path, output_path):
    """Rebase every relative external relationship without private pptx caches."""
    incoming, outgoing = io.BytesIO(blob), io.BytesIO()
    changed = []
    with ZipFile(incoming) as source, ZipFile(outgoing, "w") as dest:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.endswith(".rels"):
                root = etree.fromstring(data)
                dirty = False
                for rel in root.findall(f"{{{REL_NS}}}Relationship"):
                    if rel.get("TargetMode") != "External":
                        continue
                    old = rel.get("Target", "")
                    url = urlsplit(old)
                    if not url.path or url.scheme or url.netloc or url.path.startswith("/"):
                        continue
                    target = (input_path.parent / unquote(url.path)).resolve()
                    if not target.is_file():
                        raise ValueError(f"Relative linked file does not exist: {old}")
                    new_path = quote(Path(os.path.relpath(target, output_path.parent)).as_posix(), safe="/.-_")
                    new = urlunsplit(("", "", new_path, url.query, url.fragment))
                    rel.set("Target", new)
                    changed.append((old, new))
                    dirty = True
                if dirty:
                    data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
            dest.writestr(entry, data)
    return outgoing.getvalue(), changed


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--font", default="Apple SD Gothic Neo" if sys.platform == "darwin" else "Malgun Gothic")
    args = parser.parse_args()
    try:
        source = args.input.absolute()
        output = args.output.absolute()
        if not source.is_file() or source.suffix.lower() != ".pptx":
            raise ValueError("Input must be an existing PPTX.")
        if output.suffix.lower() != ".pptx" or output.resolve() == source.resolve():
            raise ValueError("Choose a different NEW PPTX output path; input is never overwritten.")
        if output.exists() or output.is_symlink():
            raise ValueError("Output already exists; choose a NEW output path.")
        if any(path.is_symlink() for path in output.parents):
            raise ValueError("Output directory must not use symlinks.")
        before = digest(source)
        content = load_content()
        presentation = Presentation(source)
        if len(presentation.slides) != 12:
            raise ValueError("Input must contain exactly 12 slides.")
        if abs(presentation.slide_width / presentation.slide_height - 16 / 9) > .001:
            raise ValueError("Input must be 16:9.")
        original = list(presentation.slides)
        make_intro(helpers(), presentation, content, args.font)
        slide_ids = presentation.slides._sldIdLst
        for index, slide_id in enumerate(list(slide_ids)[-4:]):
            slide_ids.remove(slide_id)
            slide_ids.insert(index, slide_id)
        for number, slide in enumerate(presentation.slides, 1):
            footers = 0
            for shape in all_shapes(slide.shapes):
                if not shape.has_text_frame:
                    continue
                if PAGE.fullmatch(shape.text.strip()) and shape.top > Inches(6.9):
                    nodes = shape._element.xpath(".//a:t")
                    nodes[0].text = f"{number:02d} / 16"
                    for node in nodes[1:]:
                        node.text = ""
                    footers += 1
            if footers != 1:
                raise ValueError(f"Slide {number}: expected exactly one page-number footer.")
            if not slide.notes_slide.notes_text_frame.text.strip():
                raise ValueError(f"Slide {number}: missing speaker notes.")
            for shape in slide.shapes:
                if (shape.left < 0 or shape.top < 0 or
                        shape.left + shape.width > presentation.slide_width + 10 or
                        shape.top + shape.height > presentation.slide_height + 10):
                    raise ValueError(f"Slide {number}: out-of-bounds shape {shape.name}")
        # Keep historical evidence text, but correct the new copy's file-location guidance.
        if output.parent.resolve() != source.parent.resolve():
            for slide in original:
                changed_location = False
                for shape in all_shapes(slide.shapes):
                    if not shape.has_text_frame:
                        continue
                    for node in shape._element.xpath(".//a:t"):
                        if node.text == "PPT와 같은 폴더에 보관":
                            node.text = "기존 MP4에 상대 링크 연결"
                            changed_location = True
                if changed_location or "영상만 PPT와 같은 폴더" in slide.notes_slide.notes_text_frame.text:
                    # Append, rather than rewriting historical source notes or their limitations.
                    slide.notes_slide.notes_text_frame.text += (
                        "\n\n소개 4장 추가 사본의 파일 위치 안내\n"
                        "위 메모의 같은 폴더 안내는 원본 덱의 위치를 기준으로 한 기록입니다. "
                        "이 사본은 기존 MP4 및 근거 파일에 대한 상대 링크를 새 출력 위치에 맞춰 조정했습니다. "
                        "사본만 이동하지 말고 원본 파일과의 상대 폴더 구조를 유지하세요. "
                        "기존 실행 근거·검토 지적·제약은 갱신하지 않았습니다."
                    )
        if len(presentation.slides) != 16:
            raise ValueError("Expected 16 slides after prepending introduction.")
        buffer = io.BytesIO()
        presentation.save(buffer)
        result, changes = rebase_links(buffer.getvalue(), source, output)
        if digest(source) != before:
            raise ValueError("Input changed during generation; refusing to write stale output.")
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as stream:
            stream.write(result)
        print(output)
        print(f"16 editable slides; 4 sourced introduction slides; input SHA-256 unchanged: {before}")
        for old, new in changes:
            print(f"Rebased: {old} -> {new}")
        return 0
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        print(f"INTRO NOT BUILT: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
